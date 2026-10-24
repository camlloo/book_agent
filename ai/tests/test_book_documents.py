import hashlib
import io
import json
import zipfile

import httpx
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from app.documents import DocumentError, parse_document, build_nodes, read_book, split_sections, Section
from app.models import BookFile, BookIndexRequest
from app.provider import Provider, ModelUnavailable


def pdf_bytes(text="PDF document fact: launch code is 5824.", encrypted=False, blank=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    if not blank:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("secret")
    output = io.BytesIO(); writer.write(output)
    return output.getvalue()


def test_text_docx_and_pdf_metadata():
    assert parse_document("中文事实".encode("utf-8-sig"), "txt")[0].text == "中文事实"
    assert parse_document("中文事实".encode("gb18030"), "md")[0].text == "中文事实"
    document = Document(); document.add_paragraph("段落甲")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "表格乙"
    document.add_paragraph("段落丙")
    output = io.BytesIO(); document.save(output)
    text = parse_document(output.getvalue(), "docx")[0].text
    assert text.index("段落甲") < text.index("表格乙") < text.index("段落丙")
    sections = parse_document(pdf_bytes(), "pdf")
    assert sections[0].page == 1 and "5824" in sections[0].text


@pytest.mark.parametrize("data,extension,message", [
    (b"bad", "pdf", "无法解析"), (b"binary\0", "txt", "二进制"),
    (pdf_bytes(encrypted=True), "pdf", "加密"), (pdf_bytes(blank=True), "pdf", "OCR"),
    (b"invalid zip", "docx", "无法解析"), (b"", "md", "没有提取")])
def test_invalid_documents_are_actionable(data, extension, message):
    with pytest.raises(DocumentError, match=message): parse_document(data, extension)


def test_docx_expansion_guard():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (61 * 1024 * 1024))
    with pytest.raises(DocumentError, match="解压后过大"): parse_document(output.getvalue(), "docx")


def test_long_chinese_and_code_chunking_is_bounded_and_does_not_need_nltk(monkeypatch):
    from llama_index.core import utils
    from llama_index.core.node_parser.text import sentence
    def unavailable(*args,**kwargs):raise PermissionError("NLTK corpus unavailable")
    monkeypatch.setattr(sentence,"split_by_sentence_tokenizer",unavailable)
    text="\n\n".join(f"第{i}段：Python 变量不需要预先声明。值是1.75，编号是{i}。\nprint('hello')" for i in range(120))
    nodes=split_sections(9,"长书稿","body","Book 正文",None,[Section(text)])
    assert len(nodes)>10
    tokenizer=utils.get_tokenizer()
    assert all(len(tokenizer(node["text"]))<=400 for node in nodes)
    joined="\n".join(node["text"] for node in nodes)
    for i in range(120): assert f"第{i}段" in joined and f"编号是{i}。" in joined


@pytest.mark.asyncio
async def test_download_integrity_origin_and_stable_chunks(monkeypatch):
    monkeypatch.setenv("MINIO_ENDPOINT", "http://minio.test:9000")
    content = "附件中的独立事实".encode()
    file = BookFile(id=8, name="证据.txt", extension="txt", size=len(content), sha256=hashlib.sha256(content).hexdigest(), url="http://minio.test:9000/private?signed=yes")
    transport = httpx.MockTransport(lambda request: httpx.Response(200, content=content))
    request = BookIndexRequest(book_id=42, revision=1, title="知识", content="手写正文", files=[file])
    nodes = await build_nodes(request, transport)
    assert {node["source_key"] for node in nodes} == {"body", "file:8"}
    assert [n["chunk_id"] for n in nodes] == [n["chunk_id"] for n in await build_nodes(request.model_copy(update={"revision":2}), transport)]
    file.url = "https://outside.test/file"
    with pytest.raises(DocumentError, match="MinIO"): await build_nodes(request, transport)
    file.url = "http://minio.test:9000/file"; file.sha256 = "0"*64
    with pytest.raises(DocumentError, match="校验值"): await build_nodes(request, transport)


@pytest.mark.asyncio
async def test_reader_uses_original_pages_without_chunk_overlap(monkeypatch):
    monkeypatch.setenv("MINIO_ENDPOINT","http://minio.test")
    data=pdf_bytes()
    file=BookFile(id=9,name="原文.pdf",extension="pdf",size=len(data),sha256=hashlib.sha256(data).hexdigest(),url="http://minio.test/file")
    result=await read_book(BookIndexRequest(book_id=7,revision=1,title="原文",content="",files=[file]),httpx.MockTransport(lambda request:httpx.Response(200,content=data)))
    assert len(result["sections"])==1
    section=result["sections"][0]
    assert section["page"]==1 and section["fileId"]==9 and section["text"].count("5824")==1


@pytest.mark.asyncio
async def test_embedding_batching_extra_body_and_dimensions(tmp_path, monkeypatch):
    monkeypatch.setenv("FIXTURE_KEY", "fixture")
    monkeypatch.setenv("EMBEDDING_PROVIDER", "remote")
    for key in ("EMBEDDING_BASE_URL", "EMBEDDING_API_KEY", "EMBEDDING_MODEL", "EMBEDDING_DIMENSIONS"):
        monkeypatch.delenv(key, raising=False)
    config = tmp_path / "config.yaml"
    config.write_text("channels:\n  - name: fixture\n    base_url: https://fixture.test\n    api_key_env: FIXTURE_KEY\n    model: chat\n    tasks: [ASK]\n    extra_body: {enable_thinking: false}\nembedding:\n  default_base_url: https://fixture.test\n  default_api_key_env: FIXTURE_KEY\n  default_model: embed\n  default_dimensions: 3\n")
    batches = []
    def handler(request):
        data = json.loads(request.content)
        if request.url.path.endswith("completions"):
            assert data["enable_thinking"] is False
            return httpx.Response(200,json={"choices":[{"message":{"content":"ok"}}]})
        batches.append(len(data["input"]))
        return httpx.Response(200,json={"data":[{"index":i,"embedding":[1,0,0]} for i in reversed(range(len(data["input"])))]})
    provider = Provider(config,httpx.MockTransport(handler))
    assert len(await provider.embeddings(["text"]*23)) == 23
    assert batches == [10,10,3]
    assert await provider.complete("ASK","system","user",history=[{"role":"user","content":"previous"},{"role":"assistant","content":"earlier answer"}]) == "ok"
    provider.transport = httpx.MockTransport(lambda request: httpx.Response(200,json={"data":[{"index":0,"embedding":[1]}]}))
    with pytest.raises(ModelUnavailable): await provider.embeddings(["text"])
